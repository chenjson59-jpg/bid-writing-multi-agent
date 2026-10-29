"""Markdown → Word 导出（核心算法 A21）。

行级解析：代码围栏、标题、表格、列表、分页符与普通段落；正文宋体 11pt、
标题黑体、表格 Table Grid。传入项目名 / 投标人时生成「封面 + 自动目录 +
页眉 + 页码页脚」的投标文件版式；不处理斜体、链接、图片、嵌套列表与脚注。
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

_HEADING_RE = re.compile(r"^(#{1,4})\s+(.*)$")
_UL_RE = re.compile(r"^[-*]\s+(.*)$")
_OL_RE = re.compile(r"^(\d+)[.、]\s+(.*)$")
_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_TOC_HEADING_RE = re.compile(r"^##\s*目录\s*$")

_HEADING_SIZES = {1: 16, 2: 14, 3: 13, 4: 12}


def markdown_to_docx(
    markdown: str,
    output_path: str | Path,
    *,
    project_name: str = "",
    bidder: str = "",
    date_text: str = "",
) -> None:
    """把 Markdown 写入 Word 文档。

    传入 `project_name` / `bidder` 时按投标文件版式导出：封面、自动目录
    （TOC 域）、页眉与「第 X 页」页脚，并去掉 Markdown 自带的标题块与手写目录；
    不传时保持纯正文转换（测试与脚本预览用）。
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    _set_default_style(doc)

    body = markdown
    if project_name or bidder:
        _add_cover(doc, project_name, bidder, date_text)
        doc.add_page_break()
        _add_toc(doc)
        doc.add_page_break()
        _apply_header_footer(doc, project_name)
        body = _body_without_front_matter(markdown)

    _render_body(doc, body)
    doc.save(str(output_path))


def _render_body(doc: Document, markdown: str) -> None:
    """行级渲染正文（A21）。"""
    table_lines: list[str] = []
    in_fence = False

    def flush_table() -> None:
        if table_lines:
            _add_table(doc, table_lines)
            table_lines.clear()

    for line in markdown.splitlines():
        stripped = line.strip()

        if stripped.startswith("```"):
            flush_table()
            in_fence = not in_fence
            continue

        if in_fence:
            doc.add_paragraph(line)
            continue

        if stripped.startswith("|"):
            table_lines.append(line)
            continue

        if table_lines:
            flush_table()

        if not stripped:
            continue

        if stripped == "---":
            doc.add_page_break()
            continue

        _add_line(doc, stripped)

    flush_table()


def _add_cover(doc: Document, project_name: str, bidder: str, date_text: str) -> None:
    """封面：项目名称、文件类型、投标人与日期（格式审查的第一页）。"""
    for _ in range(6):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(
        title.add_run(project_name or "【待填写：项目名称】"), "黑体", 22, bold=True
    )

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(subtitle.add_run("投 标 文 件"), "黑体", 26, bold=True)

    for _ in range(8):
        doc.add_paragraph()

    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(
        info.add_run(f"投标人：{bidder or '【待填写：公司名称】'}"),
        "宋体",
        14,
        bold=False,
    )

    when = doc.add_paragraph()
    when.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(when.add_run(date_text or "【待填写：日期】"), "宋体", 14, bold=False)


def _add_toc(doc: Document) -> None:
    """自动目录页：插入 TOC 域，并让 Word 打开时自动更新。"""
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(heading.add_run("目 录"), "黑体", 16, bold=True)

    paragraph = doc.add_paragraph()
    _add_field(
        paragraph,
        'TOC \\o "1-3" \\h \\z \\u',
        "（在 Word 中右键「更新域」或按 F9 生成目录）",
    )
    _enable_update_fields(doc)


def _apply_header_footer(doc: Document, project_name: str) -> None:
    """页眉写项目名，页脚居中「第 X 页」；封面首页不显示。"""
    section = doc.sections[0]
    section.different_first_page_header_footer = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_font(
        header.add_run(f"{project_name or '投标文件'} · 投标文件"),
        "宋体",
        9,
        bold=False,
    )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(footer.add_run("第 "), "宋体", 9, bold=False)
    _add_field(footer, "PAGE", "1")
    _set_font(footer.add_run(" 页"), "宋体", 9, bold=False)


def _add_field(paragraph, instruction: str, placeholder: str) -> None:
    """插入 Word 域（TOC / PAGE）：begin → instrText → separate → 占位文本 → end。"""
    run = paragraph.add_run()
    _set_font(run, "宋体", 9, bold=False)
    for tag, attrs, text in (
        ("w:fldChar", {"w:fldCharType": "begin"}, None),
        ("w:instrText", {"xml:space": "preserve"}, f" {instruction} "),
        ("w:fldChar", {"w:fldCharType": "separate"}, None),
        ("w:t", {}, placeholder),
        ("w:fldChar", {"w:fldCharType": "end"}, None),
    ):
        element = OxmlElement(tag)
        for key, value in attrs.items():
            element.set(qn(key), value)
        if text is not None:
            element.text = text
        run._r.append(element)


def _enable_update_fields(doc: Document) -> None:
    """打开文档时自动更新域，目录无需手动 F9。"""
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is not None:
        return
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def _body_without_front_matter(markdown: str) -> str:
    """去掉 Markdown 自带的标题块与手写目录（封面与自动目录已接管）。"""
    lines = markdown.splitlines()
    for index, line in enumerate(lines):
        if not _TOC_HEADING_RE.match(line.strip()):
            continue
        for inner in range(index + 1, len(lines)):
            if lines[inner].startswith("## "):
                return "\n".join(lines[inner:])
        return ""
    return markdown


def _set_default_style(doc: Document) -> None:
    """正文默认宋体 11pt（8.5.1）。"""
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def _add_line(doc: Document, line: str) -> None:
    """行级分派（8.5.2）。"""
    match = _HEADING_RE.match(line)
    if match:
        _add_heading(doc, match.group(2), len(match.group(1)))
        return

    match = _UL_RE.match(line)
    if match:
        _add_paragraph(doc, match.group(1), style="List Bullet")
        return

    match = _OL_RE.match(line)
    if match:
        _add_paragraph(doc, match.group(2), style="List Number")
        return

    _add_paragraph(doc, line)


def _add_heading(doc: Document, text: str, level: int) -> None:
    """标题使用黑体并套用内置 Heading 样式（自动目录据此取层级）。"""
    paragraph = doc.add_paragraph(style=f"Heading {min(level, 4)}")
    run = paragraph.add_run(text)
    _set_font(run, "黑体", _HEADING_SIZES.get(level, 12), bold=True)
    run.font.color.rgb = RGBColor(0, 0, 0)


def _add_paragraph(doc: Document, text: str, style: str | None = None) -> None:
    paragraph = doc.add_paragraph(style=style)
    _add_runs(paragraph, text)


def _add_runs(paragraph, text: str) -> None:
    """按粗体标记拆分 run（8.5.4）。"""
    pos = 0
    for match in _BOLD_RE.finditer(text):
        if match.start() > pos:
            _add_run(paragraph, text[pos : match.start()], bold=False)
        _add_run(paragraph, match.group(1), bold=True)
        pos = match.end()
    if pos < len(text):
        _add_run(paragraph, text[pos:], bold=False)


def _add_run(paragraph, text: str, bold: bool) -> None:
    run = paragraph.add_run(text)
    _set_font(run, "宋体", 11, bold=bold)


def _set_font(run, name: str, size: int, bold: bool) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def _add_table(doc: Document, table_lines: list[str]) -> None:
    """表格解析（8.5.3）：去分隔行、按最大列数补空。"""
    rows = [_parse_table_row(line) for line in table_lines]
    if not rows:
        return

    if len(rows) > 1 and _is_separator(rows[1]):
        rows = rows[:1] + rows[2:]
    if not rows:
        return

    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(cols):
            table.cell(i, j).text = row[j] if j < len(row) else ""


def _parse_table_row(line: str) -> list[str]:
    stripped = line.strip().removeprefix("|").removesuffix("|")
    return [cell.strip() for cell in stripped.split("|")]


def _is_separator(cells: list[str]) -> bool:
    """判断是否为 Markdown 表格分隔行（`---`、`:`、空格组成）。"""
    return bool(cells) and any("-" in cell for cell in cells) and all(
        re.fullmatch(r"[-: ]*", cell) for cell in cells
    )
