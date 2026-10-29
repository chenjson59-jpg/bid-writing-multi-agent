"""Word 导出测试（A21）：标题、段落、粗体、列表、表格、分页。"""
from __future__ import annotations

from docx import Document

from src.bid_writer.docx_export import markdown_to_docx


def _texts(path) -> list[str]:
    return [p.text for p in Document(path).paragraphs]


def test_heading_and_paragraph(tmp_path):
    out = tmp_path / "bid.docx"
    markdown_to_docx("# 标题\n\n正文内容", out)
    texts = _texts(out)
    assert "标题" in texts
    assert "正文内容" in texts


def test_table_removes_separator_row(tmp_path):
    out = tmp_path / "bid.docx"
    markdown_to_docx("| 列1 | 列2 |\n| --- | --- |\n| 甲 | 乙 |", out)
    table = Document(out).tables[0]
    assert len(table.rows) == 2
    assert table.rows[0].cells[0].text == "列1"
    assert table.rows[1].cells[0].text == "甲"


def test_bold_run(tmp_path):
    out = tmp_path / "bid.docx"
    markdown_to_docx("这是**重点**内容", out)
    runs = Document(out).paragraphs[0].runs
    assert any(run.bold for run in runs)


def test_lists(tmp_path):
    out = tmp_path / "bid.docx"
    markdown_to_docx("- 项目一\n- 项目二\n1. 第一\n2. 第二", out)
    texts = _texts(out)
    assert "项目一" in texts
    assert "第一" in texts


def test_page_break(tmp_path):
    out = tmp_path / "bid.docx"
    markdown_to_docx("上\n---\n下", out)
    texts = _texts(out)
    assert "上" in texts
    assert "下" in texts


def test_templated_cover_toc_header_footer(tmp_path):
    """投标版式（P1-3）：封面 + 自动目录域 + 页眉 + 页码页脚，去重手写标题与目录。"""
    out = tmp_path / "bid.docx"
    markdown = (
        "# 示例项目\n\n"
        "**投标人：示例公司**\n"
        "**日期：2026年10月28日**\n\n"
        "---\n\n"
        "## 目录\n1. 施工组织设计\n\n"
        "## 施工组织设计\n\n正文内容"
    )
    markdown_to_docx(
        markdown,
        out,
        project_name="示例项目",
        bidder="示例公司",
        date_text="2026年10月28日",
    )

    doc = Document(out)
    texts = _texts(out)
    assert "示例项目" in texts
    assert "投 标 文 件" in texts
    assert "投标人：示例公司" in texts
    assert "目 录" in texts
    # Markdown 自带的标题块与手写目录已被封面/自动目录接管
    assert "1. 施工组织设计" not in texts
    assert "正文内容" in texts

    # 章节使用内置 Heading 样式，TOC 域才有层级可抓
    styles = {p.text: p.style.name for p in doc.paragraphs}
    assert styles["施工组织设计"] == "Heading 2"

    # 页眉写项目名，页脚含 PAGE 域，封面首页不显示
    section = doc.sections[0]
    assert section.different_first_page_header_footer is True
    assert "示例项目" in section.header.paragraphs[0].text
    assert "PAGE" in section.footer.paragraphs[0]._p.xml

    # 目录域与「打开即更新」设置
    assert "TOC" in doc.element.body.xml
    assert "updateFields" in doc.settings.element.xml
